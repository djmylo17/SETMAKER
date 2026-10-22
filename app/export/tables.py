"""Exportar un set a Excel (.xlsx), CSV y TXT.

Todas las funciones reciben la lista de SetItem en el orden del set y el EventProfile
(para sacar el nombre de cada bloque). `genre_names` traduce el slug del género
(ej. "reggaeton_actual") a un nombre bonito (ej. "Reggaetón actual").
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable

from ..core.models import HAVE, MISSING, OTHER_VERSION, STATUS_LABELS, EventProfile, SetItem

COLUMNS = ["#", "Bloque", "Artista", "Título", "Género", "Año", "BPM", "Energía", "Estado",
           "Nuevo/Clásico", "Fuentes", "Por qué entró", "Ruta local", "Link"]

# Ancho aproximado (en caracteres) de cada columna en Excel
_WIDTHS = [5, 16, 26, 34, 18, 7, 7, 9, 13, 12, 22, 50, 50, 40]


# ---------------------------------------------------------------- ayudas
def block_name(profile: EventProfile | None, index: int) -> str:
    """Nombre del bloque; si no existe, 'Bloque N'."""
    if profile and 0 <= index < len(profile.blocks):
        return profile.blocks[index].name or f"Bloque {index + 1}"
    return f"Bloque {index + 1}" if profile and profile.blocks else ""


def genre_display(slug: str, genre_names: dict | None) -> str:
    if genre_names and slug in genre_names:
        return str(genre_names[slug])
    return slug.replace("_", " ").capitalize() if slug else ""


def first_local_path(item: SetItem) -> str:
    for lp in item.local_paths or []:
        p = lp.get("path") if isinstance(lp, dict) else str(lp)
        if p:
            return str(p)
    return ""


def all_local_paths(item: SetItem) -> list[str]:
    out = []
    for lp in item.local_paths or []:
        p = lp.get("path") if isinstance(lp, dict) else str(lp)
        if p:
            out.append(str(p))
    return out


def item_year(item: SetItem) -> int:
    return item.track.year or item.cls.year or 0


def filter_items(items: Iterable[SetItem], only_missing: bool = False) -> list[SetItem]:
    """Con only_missing=True deja solo 'Me falta' y 'Otra versión'."""
    items = list(items)
    if only_missing:
        return [i for i in items if i.status in (MISSING, OTHER_VERSION)]
    return items


def build_rows(items: Iterable[SetItem], profile: EventProfile | None, only_missing: bool = False,
               genre_names: dict | None = None) -> list[list]:
    """Filas (sin encabezado) con las columnas de COLUMNS. El # es la posición en el set completo."""
    rows = []
    for pos, it in enumerate(list(items), start=1):
        if only_missing and it.status not in (MISSING, OTHER_VERSION):
            continue
        t = it.track
        rows.append([
            pos,
            block_name(profile, it.block),
            t.artist,
            t.title,
            genre_display(it.cls.primary, genre_names),
            item_year(it) or "",
            round(t.bpm) if t.bpm else "",
            it.cls.energy if it.cls.energy else "",
            STATUS_LABELS.get(it.status, it.status),
            "Nuevo" if it.is_new else "Clásico",
            ", ".join(str(k) for k in (t.sources or {}).keys()),
            "; ".join(it.reasons or []),
            first_local_path(it),
            t.link or "",
        ])
    return rows


_BAD_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def safe_filename(name: str, max_len: int = 120) -> str:
    """Quita caracteres que Windows no permite en nombres de archivo."""
    s = _BAD_CHARS.sub(" ", name or "")
    s = re.sub(r"\s+", " ", s).strip(" .")
    if s.upper().split(".")[0] in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)),
                                   *(f"LPT{i}" for i in range(1, 10))}:
        s = "_" + s
    return (s[:max_len].rstrip(" .")) or "set"


def _prepare(path) -> Path:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


# ---------------------------------------------------------------- Excel
def export_xlsx(items: list[SetItem], profile: EventProfile | None, path, only_missing: bool = False,
                genre_names: dict | None = None) -> Path:
    """Planilla Excel lista para imprimir: encabezado oscuro, filas en negro sobre blanco."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = _prepare(path)
    wb = Workbook()
    ws = wb.active
    title = (profile.name if profile else "Set") or "Set"
    ws.title = safe_filename(("Faltantes " if only_missing else "") + title, 31).replace("[", "").replace("]", "")

    ws.append(COLUMNS)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="333333")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")

    body_font = Font(color="000000")
    link_col = COLUMNS.index("Link") + 1
    for row in build_rows(items, profile, only_missing, genre_names):
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(r, c).font = body_font
        link = row[link_col - 1]
        if link:
            cell = ws.cell(r, link_col)
            cell.hyperlink = link
            cell.font = Font(color="000000", underline="single")

    for i, w in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"
    # Impresión: horizontal y ajustada al ancho de la hoja
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    wb.save(path)
    return path


# ---------------------------------------------------------------- CSV
def export_csv(items: list[SetItem], profile: EventProfile | None, path, only_missing: bool = False,
               genre_names: dict | None = None) -> Path:
    """CSV con BOM y ';' para que Excel en Chile lo abra con tildes y columnas bien."""
    path = _prepare(path)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow(COLUMNS)
        w.writerows(build_rows(items, profile, only_missing, genre_names))
    return path


# ---------------------------------------------------------------- TXT
def export_txt(items: list[SetItem], profile: EventProfile | None, path, only_missing: bool = False,
               genre_names: dict | None = None) -> Path:
    """Texto simple: '01. Artista - Título', agrupado por bloque."""
    path = _prepare(path)
    items = list(items)
    lines: list[str] = []
    name = (profile.name if profile else "") or "Set"
    lines.append(name + (" — temas que faltan" if only_missing else ""))
    extra = " · ".join(x for x in ((profile.event_date if profile else ""),
                                   (profile.venue if profile else "")) if x)
    if extra:
        lines.append(extra)
    lines.append("")

    selected = [(pos, it) for pos, it in enumerate(items, start=1)
                if not only_missing or it.status in (MISSING, OTHER_VERSION)]
    width = max(2, len(str(len(items))))
    current = None
    for pos, it in selected:
        if it.block != current:
            current = it.block
            bname = block_name(profile, it.block)
            if bname:
                if lines and lines[-1] != "":
                    lines.append("")
                b = profile.blocks[it.block] if profile and 0 <= it.block < len(profile.blocks) else None
                head = bname + (f" ({b.start})" if b and b.start else "")
                lines.append(head)
                lines.append("-" * len(head))
        mark = ""
        if it.status == OTHER_VERSION:
            mark = "  [otra versión]"
        elif it.status == MISSING and not only_missing:
            mark = "  [me falta]"
        elif it.status == HAVE:
            mark = ""
        lines.append(f"{str(pos).zfill(width)}. {it.track.artist} - {it.track.title}{mark}")
    if not selected:
        lines.append("(No hay temas para mostrar)")
    # \r\n para que el Bloc de notas antiguo lo muestre bien
    path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8", newline="")
    return path
