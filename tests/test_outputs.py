"""Pruebas de salidas: Excel, CSV, TXT, JSON CrateDJ, links, actualizador y Spotify (sin red)."""
import csv
import io
import json
import zipfile
from urllib.parse import parse_qs, unquote, unquote_plus, urlparse

import pytest

from app.core.models import (HAVE, MISSING, OTHER_VERSION, UNKNOWN, Block, Classification,
                             EventProfile, SetItem, Track)
from app.export import (COLUMNS, build_cratedj_payload, export_cratedj_json, export_csv, export_txt,
                        export_xlsx, safe_filename, search_links)

GENRES = {"reggaeton_actual": "Reggaetón actual", "cumbia": "Cumbia"}


def _item(artist, title, block, status, genre="reggaeton_actual", paths=(), **kw):
    t = Track(key=f"{artist}|{title}".lower(), title=title, artist=artist, year=kw.get("year", 2022),
              bpm=kw.get("bpm", 96.0), isrc=kw.get("isrc", ""), deezer_id=kw.get("deezer_id", 0),
              link=kw.get("link", ""), sources={"deezer": {"rank": 1}, "apple": {}})
    return SetItem(track=t, cls=Classification(genres=[genre], energy=kw.get("energy", 7)), block=block,
                   is_new=kw.get("is_new", False), reasons=["Top en Chile", "Pedido del cliente"],
                   status=status, local_paths=[{"path": p} for p in paths])


@pytest.fixture
def profile():
    return EventProfile(name="Cumple Sofía", event_date="2026-10-25", venue="Aguas del Monte",
                        event_type="quince",
                        blocks=[Block(name="Recepción", start="21:00", minutes=60),
                                Block(name="Baile", start="22:00", minutes=120)])


@pytest.fixture
def items():
    return [
        _item("Bad Bunny", "Tití Me Preguntó", 0, HAVE, paths=[r"E:\Musica\titi.mp3", r"E:\Musica\titi2.mp3"],
              isrc="QM6MZ2214828", deezer_id=123, link="https://www.deezer.com/track/123", is_new=True),
        _item("Los Ángeles Azules", "Nunca Es Suficiente", 0, MISSING, genre="cumbia", energy=5),
        _item("Karol G", "Provenza", 1, OTHER_VERSION, paths=[r"E:\Musica\provenza remix.mp3"]),
        _item("Daddy Yankee", "Gasolina", 1, UNKNOWN, year=2004),
    ]


# ---------------------------------------------------------------- Excel
def test_xlsx(tmp_path, items, profile):
    from openpyxl import load_workbook
    path = export_xlsx(items, profile, tmp_path / "sub" / "set.xlsx", genre_names=GENRES)
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws[1][0].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.startswith("A1:")
    assert ws.max_row == 5
    row2 = [c.value for c in ws[2]]
    assert row2[1] == "Recepción" and row2[2] == "Bad Bunny" and row2[4] == "Reggaetón actual"
    assert row2[8] == "La tengo" and row2[9] == "Nuevo"
    assert row2[10] == "deezer, apple"
    assert row2[11] == "Top en Chile; Pedido del cliente"
    assert row2[12] == r"E:\Musica\titi.mp3"
    # texto negro en las filas (se imprime)
    assert ws[2][2].font.color.rgb.endswith("000000")
    assert ws.column_dimensions["D"].width > 10


def test_xlsx_only_missing(tmp_path, items, profile):
    from openpyxl import load_workbook
    ws = load_workbook(export_xlsx(items, profile, tmp_path / "f.xlsx", only_missing=True)).active
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert [r[0] for r in rows] == [2, 3]          # conserva la posición original en el set
    assert rows[0][8] == "Me falta" and rows[1][8] == "Otra versión"
    assert rows[1][1] == "Baile"


# ---------------------------------------------------------------- CSV
def test_csv(tmp_path, items, profile):
    path = export_csv(items, profile, tmp_path / "set.csv", genre_names=GENRES)
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == COLUMNS
    assert len(rows) == 5
    assert rows[2][2] == "Los Ángeles Azules" and rows[2][4] == "Cumbia"
    only = list(csv.reader(io.StringIO(export_csv(items, profile, tmp_path / "f.csv", only_missing=True)
                                       .read_text("utf-8-sig")), delimiter=";"))
    assert len(only) == 3


# ---------------------------------------------------------------- TXT
def test_txt(tmp_path, items, profile):
    text = export_txt(items, profile, tmp_path / "set.txt").read_text("utf-8")
    assert "Cumple Sofía" in text
    assert "Recepción (21:00)" in text and "Baile (22:00)" in text
    assert "01. Bad Bunny - Tití Me Preguntó" in text
    assert text.index("Recepción") < text.index("02. Los Ángeles") < text.index("Baile") < text.index("03. Karol G")
    only = export_txt(items, profile, tmp_path / "f.txt", only_missing=True).read_text("utf-8")
    assert "Bad Bunny" not in only and "02. Los Ángeles Azules - Nunca Es Suficiente" in only


def test_safe_filename():
    assert safe_filename('2026-10-25 Cumple: Sofía/Ana? - faltantes') == "2026-10-25 Cumple Sofía Ana - faltantes"
    assert safe_filename("CON") == "_CON"
    assert safe_filename("   ") == "set"


# ---------------------------------------------------------------- JSON CrateDJ
def test_cratedj_json(tmp_path, items, profile):
    path = export_cratedj_json(items, profile, tmp_path / "crate.json", genre_names=GENRES)
    data = json.loads(path.read_text("utf-8"))
    assert data["format"] == "setmaker-crate" and data["version"] == 1
    assert {"generated_at", "app_version", "event", "blocks", "tracks"} <= set(data)
    assert data["event"] == {"name": "Cumple Sofía", "date": "2026-10-25", "venue": "Aguas del Monte",
                             "event_type": "quince", "event_type_label": "15 años"}
    assert data["blocks"][1] == {"index": 1, "name": "Baile", "start": "22:00", "minutes": 120}
    keys = {"position", "block_index", "artist", "title", "genre_slug", "genre_name", "year", "bpm", "energy",
            "status", "local_path", "all_local_paths", "isrc", "deezer_id", "preview_url", "link", "sources",
            "reasons"}
    for t in data["tracks"]:
        assert set(t) == keys
        assert t["status"] in ("tengo", "otra_version", "falta")
    t0, t1, t3 = data["tracks"][0], data["tracks"][1], data["tracks"][3]
    assert t0["local_path"] == r"E:\Musica\titi.mp3" and len(t0["all_local_paths"]) == 2
    assert t0["isrc"] == "QM6MZ2214828" and t0["deezer_id"] == 123 and t0["genre_name"] == "Reggaetón actual"
    assert t1["local_path"] is None and t1["all_local_paths"] == [] and t1["isrc"] is None
    assert t3["status"] == "falta"  # sin revisar se informa como falta
    assert t0["sources"] == ["deezer", "apple"]


def test_cratedj_payload_without_profile(items):
    data = build_cratedj_payload(items, None)
    assert data["blocks"] == [] and data["tracks"][0]["genre_name"] == "Reggaeton actual"


# ---------------------------------------------------------------- links
def test_search_links_encoding():
    links = dict(search_links("Ñengo Flow & Jowell", "Hoy Se Bebe"))
    yt = urlparse(links["YouTube"])
    assert yt.netloc == "www.youtube.com"
    assert parse_qs(yt.query)["search_query"] == ["Ñengo Flow & Jowell Hoy Se Bebe"]
    assert "&" not in yt.query.split("search_query=", 1)[1]
    sp = links["Spotify"]
    assert sp.startswith("https://open.spotify.com/search/")
    assert " " not in sp and "&" not in sp
    assert unquote(sp.rsplit("/", 1)[1]) == "Ñengo Flow & Jowell Hoy Se Bebe"
    bp = urlparse(links["Beatport"])
    assert parse_qs(bp.query)["q"] == ["Ñengo Flow & Jowell Hoy Se Bebe"]
    assert len(links) == 3


def test_pool_template():
    links = search_links("Ñengo & Co", "Tití", "https://mipool.com/search?q={artista} {titulo}", "BPM Pool")
    label, url = links[-1]
    assert label == "BPM Pool"
    q = url.split("?q=", 1)[1]
    assert "&" not in q
    assert unquote_plus(q) == "Ñengo & Co Tití"
    # {query} en la ruta usa %20
    _l, url2 = search_links("A B", "C", "https://pool.cl/buscar/{query}")[-1]
    assert url2 == "https://pool.cl/buscar/A%20B%20C"
    # sin marcadores no se agrega
    assert len(search_links("a", "b", "https://pool.cl")) == 3


# ---------------------------------------------------------------- actualizador
def test_version_compare():
    from app.updater import is_newer, parse_version
    assert parse_version("v1.2.3") == ((1, 2, 3), False)
    assert is_newer("v0.2.0", "0.1.0")
    assert is_newer("0.1.10", "0.1.9")
    assert not is_newer("v0.1.0", "0.1.0")
    assert not is_newer("0.0.9", "0.1.0")
    assert is_newer("1.0.0", "1.0.0-beta.1")
    assert not is_newer("1.0.0-beta.1", "1.0.0")
    assert is_newer("1.1", "1.0.5")


def test_check_for_update_mock():
    import httpx

    from app.updater import check_for_update

    def handler(req):
        assert req.url.path == "/repos/djmylo17/setmaker/releases/latest"
        return httpx.Response(200, json={"tag_name": "v9.0.0", "body": "Novedades", "html_url": "https://gh/r",
                                         "assets": [{"name": "SETMAKER-windows.zip",
                                                     "browser_download_url": "https://gh/z.zip"}]})
    with httpx.Client(transport=httpx.MockTransport(handler)) as c:
        info = check_for_update(client=c)
    assert info == {"version": "9.0.0", "notes": "Novedades", "zip_url": "https://gh/z.zip",
                    "html_url": "https://gh/r"}
    with httpx.Client(transport=httpx.MockTransport(
            lambda r: httpx.Response(200, json={"tag_name": "v0.0.1", "assets": []}))) as c:
        assert check_for_update(client=c) is None


def test_download_and_verify(data_dir):
    import httpx

    from app.updater import UpdateError, download_update, verify_zip
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("SETMAKER/SETMAKER.exe", b"MZ")
        zf.writestr("SETMAKER/_internal/x.dll", b"x")
    payload = buf.getvalue()
    seen = []
    with httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(200, content=payload))) as c:
        path = download_update({"zip_url": "https://gh/z.zip"}, progress=lambda d, t: seen.append(d), client=c)
    assert path.name == "SETMAKER-windows.zip" and path.parent.name == "_update"
    assert seen and seen[-1] == len(payload)
    bad = path.parent / "bad.zip"
    with zipfile.ZipFile(bad, "w") as zf:
        zf.writestr("otra/cosa.txt", "x")
    with pytest.raises(UpdateError):
        verify_zip(bad)


def test_update_script_and_not_frozen(tmp_path):
    from pathlib import Path

    from app.updater import UpdateError, apply_update_and_restart, build_update_script
    bat = build_update_script(4242, Path("C:/X/data/_update/new/SETMAKER"), Path("C:/X"), "SETMAKER.exe",
                              Path("C:/X/data/_update"))
    assert 'PID eq 4242' in bat and "robocopy" in bat and "/XD data" in bat and "SETMAKER.exe" in bat
    with pytest.raises(UpdateError, match="empaquetada"):
        apply_update_and_restart(tmp_path / "x.zip")


# ---------------------------------------------------------------- Spotify (sin red)
def _sp(name, artists, uri):
    return {"name": name, "artists": [{"name": a} for a in artists], "uri": uri}


def test_spotify_pick_best_match():
    from app.sources.spotify import pick_best_match
    cands = [
        _sp("Tití Me Preguntó - Remix", ["Bad Bunny"], "spotify:track:remix"),
        _sp("Tití Me Preguntó", ["Cover Band"], "spotify:track:cover"),
        _sp("Titi Me Pregunto", ["Bad Bunny"], "spotify:track:ok"),
    ]
    assert pick_best_match("Bad Bunny", "Tití Me Preguntó (Extended Clean)", cands)["uri"] == "spotify:track:ok"
    assert pick_best_match("Karol G", "Provenza", cands) is None
    assert pick_best_match("Bad Bunny feat. X", "Algo", []) is None
    feat = [_sp("Hoy Se Bebe", ["Manuel Turizo", "Ñengo Flow"], "spotify:track:feat")]
    assert pick_best_match("Ñengo Flow", "Hoy Se Bebe", feat)["uri"] == "spotify:track:feat"


class _FakeVault:
    def __init__(self):
        self.d = {"spotify_refresh_token": "r1"}

    def get(self, k):
        return self.d.get(k, "")

    def set(self, k, v):
        if v:
            self.d[k] = v
        else:
            self.d.pop(k, None)


def test_spotify_create_playlist_mock():
    import httpx

    from app.sources.spotify import SpotifyClient
    calls = []

    def handler(req: httpx.Request):
        calls.append((req.method, req.url.path))
        if req.url.host == "accounts.spotify.com":
            return httpx.Response(200, json={"access_token": "A", "expires_in": 3600, "refresh_token": "r2"})
        if req.url.path == "/v1/search":
            assert int(req.url.params["limit"]) <= 10
            q = req.url.params["q"]
            if "Gasolina" in q:
                return httpx.Response(200, json={"tracks": {"items": [_sp("Gasolina", ["Daddy Yankee"], "u:1")]}})
            return httpx.Response(200, json={"tracks": {"items": []}})
        if req.url.path == "/v1/me/playlists":
            return httpx.Response(201, json={"id": "PL", "external_urls": {"spotify": "https://open.spotify.com/playlist/PL"}})
        if req.url.path == "/v1/playlists/PL/items":
            assert json.loads(req.content)["uris"] == ["u:1"]
            return httpx.Response(201, json={"snapshot_id": "s"})
        return httpx.Response(404)

    vault = _FakeVault()
    client = SpotifyClient(vault, "a" * 32, http=httpx.Client(transport=httpx.MockTransport(handler)))
    its = [_item("Daddy Yankee", "Gasolina", 0, MISSING), _item("Nadie", "No Existe", 0, MISSING)]
    prog = []
    res = client.create_playlist_from_items(its, "Test", progress=lambda n, t: prog.append((n, t)))
    assert res == {"url": "https://open.spotify.com/playlist/PL", "added": 1, "not_found": ["Nadie - No Existe"]}
    assert vault.d["spotify_refresh_token"] == "r2"     # se guarda el token rotado
    assert ("POST", "/v1/me/playlists") in calls and ("POST", "/v1/playlists/PL/items") in calls
    assert prog[-1] == (2, 2)


def test_spotify_403_message():
    import httpx

    from app.net import SourceError
    from app.sources.spotify import SpotifyClient

    def handler(req):
        if req.url.host == "accounts.spotify.com":
            return httpx.Response(200, json={"access_token": "A", "expires_in": 3600})
        return httpx.Response(403)
    client = SpotifyClient(_FakeVault(), "a" * 32, http=httpx.Client(transport=httpx.MockTransport(handler)))
    with pytest.raises(SourceError, match="permiso"):
        client.search_track("A", "B")


def test_spotify_pkce_and_auth_url():
    from app.sources.spotify import SpotifyClient, _pkce_pair
    v, c = _pkce_pair()
    assert 43 <= len(v) <= 128 and "=" not in c
    url = SpotifyClient(_FakeVault(), "abc").build_auth_url("st", c)
    q = parse_qs(urlparse(url).query)
    assert q["redirect_uri"] == ["http://127.0.0.1:8765/callback"]
    assert q["code_challenge_method"] == ["S256"] and q["state"] == ["st"]
    assert q["scope"] == ["playlist-modify-private playlist-modify-public"]


# ---------------------------------------------------------------- UI (offscreen)
@pytest.fixture(scope="module")
def qapp():
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def test_ui_modules(qapp, data_dir, items, profile):
    from app.main import build_context
    from app.ui.pages.missing_page import MissingPage
    from app.ui.pages.outputs_cards import OutputsCard, UpdatesCard

    ctx = build_context()
    ctx.vault.unlock()
    msgs = []
    page = MissingPage(ctx, msgs.append)
    assert not page.b_full.isEnabled()
    ctx.config.set("links.pool_url", "https://mipool.com/search?q={artista} {titulo}")
    page.set_set(profile, items, GENRES)
    assert page.table.rowCount() == 2
    assert page.b_full.isEnabled()
    btns = page.table.cellWidget(0, 6).findChildren(type(page.b_full))
    assert [b.text() for b in btns] == ["YouTube", "Spotify", "Beatport", "Mi pool"]
    assert page.file_name("faltantes", "xlsx") == "2026-10-25 Cumple Sofía - faltantes.xlsx"
    assert page.export_dir() == ctx.data_dir / "exports"
    page.set_set(None, [], {})
    assert page.table.rowCount() == 0
    OutputsCard(ctx, msgs.append)
    UpdatesCard(ctx, msgs.append)
    ctx.db.close()
